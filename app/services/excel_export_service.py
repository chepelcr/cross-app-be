import logging
from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.enums.report_color import get_color_palette
from app.models.order import Order
from app.services.pdf_service import _s3_key, upload_file_to_s3

logger = logging.getLogger(__name__)

COLUMNS = [
    ("PO", 15),
    ("UPC", 18),
    ("LINE_NBR", 10),
    ("ITEM", 15),
    ("STORE_NBR", 12),
    ("ORDENADO", 12),
    ("CANTIDAD A ENTREGAR", 22),
    ("DIFERENCIA", 12),
]

HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def generate_nuevo_reporte(order: Order, crossdocking_data, color=None) -> bytes:
    """Generate NuevoReporte Excel workbook from crossdocking data."""
    palette = get_color_palette(color)
    header_fill = PatternFill(
        start_color=palette["excel_header"],
        end_color=palette["excel_header"],
        fill_type="solid",
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NuevoReporte"

    # Write header row
    for col_idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = HEADER_FONT
        cell.fill = header_fill
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    # Write data rows: one row per item per sale point
    row_num = 2
    for sp in crossdocking_data.sale_points:
        line_nbr = 0
        for item in sp.items:
            line_nbr += 1
            ordered = item.quantity or 0
            sent = item.sent or 0
            diferencia = ordered - sent

            values = [
                order.document_number or "",
                item.original_code or "",
                line_nbr,
                item.internal_code or "",
                sp.store_number or "",
                ordered,
                sent,
                diferencia,
            ]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=row_num, column=col_idx, value=val)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="center")
            row_num += 1

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def create_nuevo_reporte(order: Order, crossdocking_data, color=None) -> str:
    """Generate NuevoReporte Excel and upload to S3."""
    excel_bytes = generate_nuevo_reporte(order, crossdocking_data, color)
    last4 = (order.document_number or "")[-4:]
    key = _s3_key(order.company_id, order.document_number, f"{last4}-RN.xlsx")
    url = upload_file_to_s3(
        excel_bytes,
        key,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    return url
