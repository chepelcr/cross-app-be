from __future__ import annotations

import logging
from io import BytesIO
from typing import Dict, List, Optional
import uuid

import openpyxl

from app.dtos import ExcelFileDTO
from app.dtos.requests.store_request_dto import StoreRequestDTO
from app.dtos.responses.pagination_dto import PaginationResponse
from app.dtos.responses.store_dto import StoreListResponse, StoreResponse
from app.enums.store_search_filters import StoreSearchFilters
from app.models.store import Store
from app.repositories.store_repository import StoreRepository
from app.utils.crossdocking_utils import decode_excel_file
from app.utils.search_utils import SearchUtils

logger = logging.getLogger(__name__)


def get_stores(
    company_id: str,
    client_id: str,
    page: int = 1,
    page_size: int = 12,
    search: str = None,
) -> StoreListResponse:
    """Get paginated stores for a company/client with optional search filters."""
    search_filters = None
    order_by = None

    if search:
        filters, order_result = SearchUtils.parse_search_filter(
            search, Store, StoreSearchFilters
        )
        if filters:
            search_filters = filters
        if order_result:
            order_by = order_result

    client_uuid = uuid.UUID(client_id)

    with StoreRepository() as repo:
        stores, total = repo.find_all_by_client(
            company_id,
            client_uuid,
            search_filters=search_filters,
            order_by=order_by,
            page=page,
            page_size=page_size,
        )

    total_pages = (total + page_size - 1) // page_size if total > 0 else 1
    return StoreListResponse(
        data=[_map_store(s) for s in stores],
        pagination=PaginationResponse(
            page=page,
            pageSize=page_size,
            totalElements=total,
            totalPages=total_pages,
        ),
    )


def get_store(company_id: str, store_id: uuid.UUID) -> Optional[StoreResponse]:
    """Get a single store by ID."""
    with StoreRepository() as repo:
        store = repo.find_by_id_and_company(store_id, company_id)
    if not store:
        return None
    return _map_store(store)


def create_store(
    company_id: str,
    client_id_str: str,
    dto: "StoreRequestDTO",
) -> StoreResponse:
    """Create a new store."""
    client_id = uuid.UUID(client_id_str)

    with StoreRepository() as repo:
        store = Store(
            company_id=company_id,
            client_id=client_id,
            store_code=dto.store_code,
            store_name=dto.store_name,
            slot_id=dto.slot_id,
            chain=dto.chain,
        )
        store = repo.save(store)

    return _map_store(store)


def update_store(
    company_id: str,
    store_id_str: str,
    dto: "StoreRequestDTO",
) -> Optional[StoreResponse]:
    """Update an existing store."""
    store_id = uuid.UUID(store_id_str)

    with StoreRepository() as repo:
        store = repo.find_by_id_and_company(store_id, company_id)
        if not store:
            return None

        if dto.store_code is not None:
            store.store_code = dto.store_code
        if dto.store_name is not None:
            store.store_name = dto.store_name
        if dto.slot_id is not None:
            store.slot_id = dto.slot_id
        if dto.chain is not None:
            store.chain = dto.chain

        store = repo.save(store)

    return _map_store(store)


def update_store_status(
    company_id: str,
    store_id_str: str,
    status: int,
) -> Optional[StoreResponse]:
    """Update a store's status."""
    store_id = uuid.UUID(store_id_str)

    with StoreRepository() as repo:
        store = repo.find_by_id_and_company(store_id, company_id)
        if not store:
            return None

        store.status = status
        store = repo.save(store)

    return _map_store(store)


def upload_stores_excel(company_id: str, client_id_str: str, body: ExcelFileDTO) -> int:
    """Upload stores from an Excel file. Returns count of upserted records."""
    file = decode_excel_file(body)
    client_id = uuid.UUID(client_id_str)

    wb = openpyxl.load_workbook(file, read_only=True, data_only=True, keep_vba=True)
    ws = wb["PUNTOS DE VENTA"] if "PUNTOS DE VENTA" in wb.sheetnames else wb.active

    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        wb.close()
        return 0

    # Find header row
    header_row = rows[0]
    headers = [str(h).strip().lower() if h else "" for h in header_row]

    col_map = {}
    for idx, header in enumerate(headers):
        if header in ("codigo", "código"):
            col_map["store_code"] = idx
        elif header == "nombre":
            col_map["store_name"] = idx
        elif header in ("slot id", "slot_id", "slotid"):
            col_map["slot_id"] = idx
        elif header == "cadena":
            col_map["chain"] = idx

    if "store_code" not in col_map:
        wb.close()
        raise ValueError("Missing required column: Codigo")

    stores_data: List[dict] = []
    for row in rows[1:]:
        code_val = row[col_map["store_code"]] if col_map.get("store_code") is not None else None
        if not code_val:
            continue

        store_entry = {
            "store_code": str(code_val).strip(),
        }
        if "store_name" in col_map and row[col_map["store_name"]]:
            store_entry["store_name"] = str(row[col_map["store_name"]]).strip()
        if "slot_id" in col_map and row[col_map["slot_id"]]:
            store_entry["slot_id"] = str(row[col_map["slot_id"]]).strip()
        if "chain" in col_map and row[col_map["chain"]]:
            store_entry["chain"] = str(row[col_map["chain"]]).strip()

        stores_data.append(store_entry)

    wb.close()

    if not stores_data:
        return 0

    with StoreRepository() as repo:
        count = repo.bulk_upsert(company_id, client_id, stores_data)

    return count


def get_slot_map(company_id: str, client_id: str) -> Dict[str, str]:
    """Get store_code -> slot_id mapping for a company/client."""
    client_uuid = uuid.UUID(client_id)
    with StoreRepository() as repo:
        return repo.get_slot_map(company_id, client_uuid)


def _map_store(store: Store) -> StoreResponse:
    return StoreResponse(
        storeId=str(store.store_id),
        companyId=store.company_id,
        clientId=str(store.client_id),
        storeCode=store.store_code,
        storeName=store.store_name,
        slotId=store.slot_id,
        chain=store.chain,
        gln=store.gln,
    )
